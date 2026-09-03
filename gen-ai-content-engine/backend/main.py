import os
import re
import secrets
import hashlib
import hmac
import base64
import smtplib
import json
import time
import uuid
import threading
import sqlite3
import tempfile
import zipfile
from email.message import EmailMessage
from typing import Optional
from pathlib import Path
from dotenv import load_dotenv

# Load .env from the workspace root (two levels up from backend/)
_env_path = Path(__file__).parent.parent.parent / ".env"
load_dotenv(dotenv_path=_env_path)
load_dotenv()  # also try local as fallback

from fastapi import FastAPI, UploadFile, Form, HTTPException, Request, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from starlette.concurrency import run_in_threadpool
import pypdf
import pytesseract
from PIL import Image
import io
from groq import Groq
from mem0 import MemoryClient

app = FastAPI()
AUTH_SECRET = os.environ.get("AUTH_SECRET")
DEV_AUTH_FALLBACK = os.environ.get("DEV_AUTH_FALLBACK", "false").lower() == "true"
FRONTEND_ORIGINS = [x.strip() for x in os.environ.get("FRONTEND_ORIGINS", "http://localhost:5173,http://localhost:3000").split(",") if x.strip()]
if not AUTH_SECRET and not DEV_AUTH_FALLBACK:
    raise RuntimeError("AUTH_SECRET must be configured")

app.add_middleware(
    CORSMiddleware,
    allow_origins=FRONTEND_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Development-only store. Replace this repository with Redis for multi-worker production deployments.
otp_store = {}
OTP_EXPIRY_SECONDS = 300  # 5 minutes
OTP_MAX_ATTEMPTS = 5
RESEND_COOLDOWN_SECONDS = 60
MAX_HISTORY_ENTRIES = 50
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
SUPPORTED_OUTPUTS = {
    "Exec Summary", "Advisory", "LinkedIn Post", "Video Script", "Presentation",
    "Twitter/X Thread", "Infographic", "Press Release", "Email Brief", "FAQ",
}
SUPPORTED_TONES = {
    "Professional", "Authoritative & Strategic", "Casual & Engaging",
    "Urgent & Action-Oriented", "Inspirational",
}
SUPPORTED_AUDIENCES = {
    "Leadership / Execs", "General Public", "Tech / Developers", "Sales / Marketing",
    "Stakeholders & Investors",
}
SUPPORTED_LANGUAGES = {"English (US)", "English (UK)", "Hindi", "Spanish", "French", "Arabic"}
SUPPORTED_DETAIL_LEVELS = {"Concise", "Standard", "Comprehensive"}
SUPPORTED_OBJECTIVES = {"Inform", "Persuade", "Educate", "Alert", "Drive action"}
SUPPORTED_STYLES = {"Clear & direct", "Executive briefing", "Story-led", "Data-led", "Conversational"}

# Note: Ensure groq is installed. Tesseract OS binaries must also be installed (e.g., apt-get install tesseract-ocr)
groq_key = os.environ.get("GROQ_API_KEY")
if not groq_key:
    raise RuntimeError("GROQ_API_KEY must be configured")
groq_client = Groq(api_key=groq_key, timeout=30.0, max_retries=2)
mem0_client = MemoryClient(api_key=os.environ["MEM0_API_KEY"]) if os.environ.get("MEM0_API_KEY") else None
HISTORY_FILE = Path(__file__).with_name("chat_history.json")
ACCOUNT_DB = Path(__file__).with_name("accounts.sqlite3")
history_lock = threading.Lock()

def _account_db():
    db = sqlite3.connect(ACCOUNT_DB)
    db.execute("CREATE TABLE IF NOT EXISTS accounts (email TEXT PRIMARY KEY, salt TEXT NOT NULL, password_hash TEXT NOT NULL, updated_at INTEGER NOT NULL)")
    return db

def _password_hash(password, salt=None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 240_000).hex()
    return salt, digest

def _memories(user_id):
    # Structured account/history records must be read exhaustively; semantic
    # search can omit an otherwise valid JSON record from its ranked results.
    if not mem0_client:
        return []
    memories = []
    page = 1
    while True:
        result = mem0_client.get_all(filters={"user_id": user_id}, page=page, page_size=100)
        if isinstance(result, dict):
            memories.extend(result.get("results", []))
            if not result.get("next"):
                break
            page += 1
        else:
            memories.extend(result or [])
            break
    return memories

def _memory_payload(memory):
    """Return the JSON history payload embedded in a Mem0 memory, if present."""
    if not isinstance(memory, dict):
        return None
    for key in ("memory", "text", "content"):
        value = memory.get(key)
        if isinstance(value, dict):
            return value
        if isinstance(value, str):
            try:
                payload = json.loads(value)
            except (TypeError, ValueError):
                continue
            if isinstance(payload, dict):
                return payload
    return None

def _delete_memory_copy(user_id, history_id):
    """Delete all Mem0 records carrying this local history ID."""
    if not mem0_client:
        return
    for memory in _memories(user_id):
        payload = _memory_payload(memory)
        if not payload or str(payload.get("id")) != history_id:
            continue
        memory_id = memory.get("id") or memory.get("memory_id")
        if memory_id:
            mem0_client.delete(memory_id=memory_id)

def _account(email):
    with _account_db() as db:
        row = db.execute("SELECT email, salt, password_hash, updated_at FROM accounts WHERE email = ?", (email,)).fetchone()
    return dict(zip(("email", "salt", "password_hash", "updated_at"), row)) if row else None

def _save_memory(email, payload):
    if mem0_client:
        mem0_client.add(messages=[{"role": "user", "content": json.dumps(payload)}], user_id=email, infer=False)

def _save_account(email, password):
    salt, password_hash = _password_hash(password)
    with _account_db() as db:
        db.execute("INSERT INTO accounts(email, salt, password_hash, updated_at) VALUES (?, ?, ?, ?) ON CONFLICT(email) DO UPDATE SET salt=excluded.salt, password_hash=excluded.password_hash, updated_at=excluded.updated_at", (email, salt, password_hash, int(time.time())))

def _save_exact_history(email, payload):
    with history_lock:
        try:
            records = json.loads(HISTORY_FILE.read_text(encoding="utf-8")) if HISTORY_FILE.exists() else {}
        except (OSError, ValueError):
            records = {}
        user_records = records.setdefault(email, [])
        user_records.append(payload)
        records[email] = user_records[-MAX_HISTORY_ENTRIES:]
        HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(prefix="chat_history.", suffix=".tmp", dir=HISTORY_FILE.parent)
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as temp_file:
                json.dump(records, temp_file, ensure_ascii=False)
                temp_file.flush()
                os.fsync(temp_file.fileno())
            os.replace(temp_name, HISTORY_FILE)
        finally:
            if os.path.exists(temp_name):
                os.unlink(temp_name)

def _read_exact_history(email):
    with history_lock:
        try:
            records = json.loads(HISTORY_FILE.read_text(encoding="utf-8")) if HISTORY_FILE.exists() else {}
            return records.get(email, [])
        except (OSError, ValueError):
            return []

def _otp_hash(email, code):
    return hmac.new((AUTH_SECRET or "dev-secret").encode(), f"{email}:{code}".encode(), hashlib.sha256).hexdigest()

def _session_token(email):
    payload = f"{email}:{int(time.time()) + 3600}".encode()
    encoded = base64.urlsafe_b64encode(payload).decode().rstrip("=")
    sig = hmac.new((AUTH_SECRET or "dev-secret").encode(), encoded.encode(), hashlib.sha256).hexdigest()
    return f"{encoded}.{sig}"

def _authenticated(request: Request):
    token = request.cookies.get("session")
    if not token or "." not in token: return False
    encoded, signature = token.rsplit(".", 1)
    if not hmac.compare_digest(signature, hmac.new((AUTH_SECRET or "dev-secret").encode(), encoded.encode(), hashlib.sha256).hexdigest()): return False
    try:
        email, expiry = base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4)).decode().rsplit(":", 1)
        return bool(email and time.time() < int(expiry))
    except (ValueError, TypeError): return False


def _generation_option(value: Optional[str], allowed: set[str], label: str) -> str:
    """Validate operator-controlled generation settings at the API boundary."""
    normalized = str(value or "").strip()
    if normalized not in allowed:
        raise HTTPException(status_code=400, detail=f"Unsupported {label}.")
    return normalized

# Configure Resend if key exists
resend_api_key = os.environ.get("RESEND_API_KEY")
if resend_api_key:
    try:
        import resend
        resend.api_key = resend_api_key
    except Exception as e:
        print("Resend import error:", e)

class AuthReq(BaseModel):
    email: str
    password: Optional[str] = None
    code: Optional[str] = None
    new_password: Optional[str] = None

def dispatch_email(recipient_email: str, otp_code: str) -> dict:
    """Dispatches the OTP email via Resend API, Gmail SMTP, or logs to terminal."""
    delivery_method = "terminal_fallback"
    
    # 1. Try Resend API (Recommended for production)
    if os.environ.get("RESEND_API_KEY"):
        try:
            import resend
            from_sender = os.environ.get("RESEND_FROM_EMAIL", "onboarding@resend.dev")
            resend.Emails.send({
                "from": f"Gen AI Engine <{from_sender}>",
                "to": [recipient_email],
                "subject": f"Your Verification Code: {otp_code}",
                "text": f"Your verification code is: {otp_code}\nThis code will expire in 5 minutes.",
                "html": f"""
                <div style="font-family: Arial, sans-serif; padding: 20px; color: #333;">
                    <h2>Gen AI Transformation Engine</h2>
                    <p>Your one-time login verification code is:</p>
                    <div style="font-size: 28px; font-weight: bold; letter-spacing: 4px; color: #005C4B; padding: 12px 0;">
                        {otp_code}
                    </div>
                    <p style="color: #666; font-size: 13px;">This code is valid for 5 minutes. If you didn't request this, you can ignore this email.</p>
                </div>
                """
            })
            delivery_method = "resend"
            print(f"[EMAIL] Successfully sent OTP via Resend to {recipient_email}")
            return {"status": "sent", "channel": "resend"}
        except Exception as e:
            print("[EMAIL] Resend delivery failed:", e)

    # 2. Try Gmail SMTP
    smtp_email = os.environ.get("SMTP_EMAIL")
    # Google displays app passwords with spaces; SMTP expects the 16-character
    # value without formatting whitespace.
    smtp_password = re.sub(r"\s+", "", os.environ.get("SMTP_APP_PASSWORD", ""))
    if smtp_email and smtp_password:
        try:
            msg = EmailMessage()
            msg.set_content(f"Your verification code is: {otp_code}\nValid for 5 minutes.")
            msg["Subject"] = f"Your Verification Code: {otp_code}"
            msg["From"] = smtp_email
            msg["To"] = recipient_email
            
            server = smtplib.SMTP("smtp.gmail.com", 587)
            server.starttls()
            server.login(smtp_email, smtp_password)
            server.send_message(msg)
            server.quit()
            delivery_method = "smtp"
            print(f"[EMAIL] Successfully sent OTP via Gmail SMTP to {recipient_email}")
            return {"status": "sent", "channel": "smtp"}
        except Exception as e:
            print("[EMAIL] SMTP delivery failed:", e)
    
    if not DEV_AUTH_FALLBACK:
        return {"status": "unavailable", "channel": "email"}
    # 3. Explicit development-only console fallback
    print(f"\n=======================================================")
    print(f" [LOCAL OTP] Code for {recipient_email} : {otp_code}")
    print(f" (Copy this code to log in immediately)")
    print(f"=======================================================\n")
    return {"status": "sent", "channel": "console"}

@app.post("/auth/send")
@app.post("/api/auth/send")
def send_otp(req: AuthReq, request: Request, response: Response):
    if not req.email or "@" not in req.email:
        raise HTTPException(status_code=400, detail="Please enter a valid email address.")
        
    email = req.email.lower().strip()
    if not req.password or len(req.password) < 8:
        raise HTTPException(status_code=400, detail="Password is required and must be at least 8 characters.")
    account = _account(email)
    if account:
        _, candidate = _password_hash(req.password, account["salt"])
        if not hmac.compare_digest(candidate, account["password_hash"]):
            raise HTTPException(status_code=401, detail="Invalid email or password.")
        response.set_cookie("session", _session_token(email), httponly=True, secure=not DEV_AUTH_FALLBACK, samesite="lax", max_age=3600)
        return {"authenticated": True, "email": email, "message": "Signed in with your saved password."}
    now = time.time()
    for key in list(otp_store):
        if otp_store[key].get("expires_at", 0) <= now: del otp_store[key]
    ip_key = f"ip:{request.client.host if request.client else 'unknown'}"
    previous = otp_store.get(email)
    ip_previous = otp_store.get(ip_key)
    if ((previous and now - previous.get("sent_at", 0) < RESEND_COOLDOWN_SECONDS) or
            (ip_previous and now - ip_previous.get("sent_at", 0) < RESEND_COOLDOWN_SECONDS)):
        raise HTTPException(status_code=429, detail="Please wait before requesting another code.")
    otp = f"{secrets.randbelow(1_000_000):06d}"
    otp_store[email] = {
        "hash": _otp_hash(email, otp),
        "expires_at": time.time() + OTP_EXPIRY_SECONDS
        , "attempts": 0, "sent_at": now
    }
    otp_store[ip_key] = {"sent_at": now, "expires_at": now + RESEND_COOLDOWN_SECONDS}
    delivery = dispatch_email(email, otp)
    if delivery.get("status") != "sent":
        otp_store.pop(email, None)
        otp_store.pop(ip_key, None)
        raise HTTPException(status_code=503, detail="Verification email could not be sent. Please try again later.")
    return {"status": "sent", "message": "If the address is eligible, a verification code has been sent."}

@app.post("/auth/verify")
@app.post("/api/auth/verify")
def verify_otp(req: AuthReq, response: Response):
    email = req.email.lower().strip() if req.email else ""
    if not req.password or len(req.password) < 8:
        raise HTTPException(status_code=400, detail="Password is required and must be at least 8 characters.")
    user_entry = otp_store.get(email)
    
    if not user_entry:
        raise HTTPException(status_code=400, detail="Invalid or expired verification code.")
        
    # Check expiry (5 min)
    if time.time() > user_entry["expires_at"]:
        del otp_store[email]
        raise HTTPException(status_code=400, detail="Invalid or expired verification code.")
        
    user_entry["attempts"] += 1
    valid = hmac.compare_digest(user_entry["hash"], _otp_hash(email, str(req.code or "").strip()))
    if valid:
        del otp_store[email]
        _save_account(email, req.password)
        response.set_cookie("session", _session_token(email), httponly=True, secure=not DEV_AUTH_FALLBACK, samesite="lax", max_age=3600)
        return {"authenticated": True, "email": email}
    if user_entry["attempts"] >= OTP_MAX_ATTEMPTS: del otp_store[email]
    raise HTTPException(status_code=400, detail="Invalid or expired verification code.")

@app.post("/auth/forgot/send")
@app.post("/api/auth/forgot/send")
def forgot_send(req: AuthReq, request: Request):
    email = req.email.lower().strip() if req.email else ""
    if not email or "@" not in email: raise HTTPException(status_code=400, detail="Please enter a valid email address.")
    now = time.time()
    key = f"forgot:{email}"
    ip_key = f"forgot-ip:{request.client.host if request.client else 'unknown'}"
    previous = otp_store.get(key)
    ip_previous = otp_store.get(ip_key)
    if ((previous and now - previous.get("sent_at", 0) < RESEND_COOLDOWN_SECONDS) or
            (ip_previous and now - ip_previous.get("sent_at", 0) < RESEND_COOLDOWN_SECONDS)):
        raise HTTPException(status_code=429, detail="Please wait before requesting another code.")
    otp = f"{secrets.randbelow(1_000_000):06d}"
    otp_store[key] = {"hash": _otp_hash(email, otp), "expires_at": now + OTP_EXPIRY_SECONDS, "attempts": 0, "sent_at": now}
    otp_store[ip_key] = {"sent_at": now, "expires_at": now + RESEND_COOLDOWN_SECONDS}
    delivery = dispatch_email(email, otp)
    if delivery.get("status") != "sent":
        otp_store.pop(key, None)
        otp_store.pop(ip_key, None)
        raise HTTPException(status_code=503, detail="Verification email could not be sent. Please try again later.")
    return {"status": "sent", "channel": delivery.get("channel"), "message": "If the address is eligible, a verification code has been sent."}

@app.post("/auth/forgot/verify")
@app.post("/api/auth/forgot/verify")
def forgot_verify(req: AuthReq, response: Response):
    email = req.email.lower().strip() if req.email else ""
    if not _account(email):
        raise HTTPException(status_code=400, detail="Invalid or expired verification code.")
    entry = otp_store.get(f"forgot:{email}")
    if not entry or time.time() > entry["expires_at"]: raise HTTPException(status_code=400, detail="Invalid or expired verification code.")
    if entry["attempts"] >= OTP_MAX_ATTEMPTS:
        del otp_store[f"forgot:{email}"]
        raise HTTPException(status_code=400, detail="Invalid or expired verification code.")
    if not hmac.compare_digest(entry["hash"], _otp_hash(email, str(req.code or "").strip())):
        entry["attempts"] += 1
        if entry["attempts"] >= OTP_MAX_ATTEMPTS:
            del otp_store[f"forgot:{email}"]
        raise HTTPException(status_code=400, detail="Invalid or expired verification code.")
    if not req.new_password or len(req.new_password) < 8: raise HTTPException(status_code=400, detail="New password must be at least 8 characters.")
    del otp_store[f"forgot:{email}"]
    _save_account(email, req.new_password)
    response.set_cookie("session", _session_token(email), httponly=True, secure=not DEV_AUTH_FALLBACK, samesite="lax", max_age=3600)
    return {"authenticated": True, "email": email, "message": "Password reset successfully."}

@app.post("/auth/logout")
@app.post("/api/auth/logout")
def logout(response: Response):
    response.delete_cookie("session")
    return {"authenticated": False}

def _extract_file_text(content, filename):
    """Run CPU-bound document extraction away from the async request loop."""
    if filename.endswith((".txt", ".md")):
        return content.decode("utf-8", errors="ignore")
    if filename.endswith(".docx"):
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX support is unavailable. Install python-docx and restart the API.") from exc
        try:
            document = Document(io.BytesIO(content))
            return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        except (KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
            raise ValueError("The DOCX file could not be read.") from exc
    if filename.endswith(".pdf"):
        try:
            pdf = pypdf.PdfReader(io.BytesIO(content))
            return "\n".join(filter(None, (page.extract_text() for page in pdf.pages)))
        except (ValueError, pypdf.errors.PdfReadError, OSError) as exc:
            raise ValueError("The PDF could not be read.") from exc
    if filename.endswith((".png", ".jpg", ".jpeg")):
        try:
            image = Image.open(io.BytesIO(content))
            image.verify()
            image = Image.open(io.BytesIO(content))
            return pytesseract.image_to_string(image)
        except (OSError, ValueError) as exc:
            raise ValueError("The image could not be read.") from exc
    raise ValueError("Unsupported file type.")

@app.post("/transform")
@app.post("/api/transform")
async def transform_content(
    request: Request,
    file: Optional[UploadFile] = None,
    text: Optional[str] = Form(""),
    outputs: Optional[str] = Form("[]"),
    tone: Optional[str] = Form("Professional"),
    audience: Optional[str] = Form("Leadership / Execs"),
    language: Optional[str] = Form("English (US)"),
    detail_level: Optional[str] = Form("Standard"),
    objective: Optional[str] = Form("Inform"),
    style: Optional[str] = Form("Clear & direct"),
    source_fidelity: Optional[str] = Form("true"),
):
    if not _authenticated(request): raise HTTPException(status_code=401, detail="Authentication required.")
    try:
        out_list = json.loads(outputs or "")
    except (TypeError, json.JSONDecodeError):
        raise HTTPException(status_code=400, detail="outputs must be valid JSON.")
    if not isinstance(out_list, list) or not out_list: raise HTTPException(status_code=400, detail="outputs must be a non-empty JSON array.")
    if any(not isinstance(item, str) or item not in SUPPORTED_OUTPUTS for item in out_list): raise HTTPException(status_code=400, detail="Unsupported output format.")
    tone = _generation_option(tone, SUPPORTED_TONES, "tone")
    audience = _generation_option(audience, SUPPORTED_AUDIENCES, "target audience")
    language = _generation_option(language, SUPPORTED_LANGUAGES, "language")
    detail_level = _generation_option(detail_level, SUPPORTED_DETAIL_LEVELS, "detail level")
    objective = _generation_option(objective, SUPPORTED_OBJECTIVES, "communication objective")
    style = _generation_option(style, SUPPORTED_STYLES, "content style")
    source_fidelity = str(source_fidelity or "true").lower() in {"1", "true", "yes", "on"}
    extracted_text = ""
    if file:
        content = await file.read()
        if len(content) > MAX_UPLOAD_BYTES: raise HTTPException(status_code=400, detail="File exceeds the 10 MB limit.")
        filename = (file.filename or "").lower()
        if not filename: raise HTTPException(status_code=400, detail="A filename is required.")
        try:
            extracted_text = await run_in_threadpool(_extract_file_text, content, filename)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
            
    combined_text = f"{text}\n\n{extracted_text}".strip()
    
    if not combined_text:
        raise HTTPException(status_code=400, detail="No input text or document provided.")
        
    combined_text = combined_text[:100_000]
    
    system_prompt = """You are an elite, world-class Content Strategist and Executive Communications Specialist.
Your objective is to transform the user's source material into production-ready deliverables for each requested format.

CRITICAL CONTENT QUALITY RULES:
1. MATCH THE OPERATOR'S BRIEF: Honor the requested audience, tone, language, detail level, communication objective, and content style. Make every deliverable complete for its format without adding filler.
2. TAILOR ACCORDING TO FORMAT:
   - 'Exec Summary': Comprehensive executive briefing (Context & Background, Core Findings & Metrics, Strategic Implications, Risk Assessment, and 4-6 Detailed Actionable Recommendations).
   - 'Advisory': Formal strategic advisory bulletin (Executive Summary, Threat/Opportunity Matrix, Core Strategic Imperatives, Governance/Tech Guidance, and Step-by-Step Implementation Roadmap).
   - 'LinkedIn Post': High-engagement professional structure (Hook with spacing, storytelling context, clear bullet insights, key lessons, thought-provoking closing question, and relevant industry hashtags). Never use emojis.
   - 'Video Script': Complete video package: creative brief, storyboard with scene number/duration/scene description, visual and B-roll recommendations, on-screen text, verbatim narration, subtitle-ready dialogue, sound/music cues, and a strong CTA.
   - 'Presentation': Slide-by-slide deck outline (Slide Number, Slide Title, Key Bullet Points, Visual/Graphic Concept, and Verbatim Speaker Script).
   - 'Twitter/X Thread': 5-8 tweet thread numbered (1/X), hook tweet, body breakdown, and recap CTA tweet.
   - 'Infographic': A structured visual-first content breakdown — a punchy Title, 5-7 bold Headline Stats or Key Facts (each with a short 1-sentence explanation), 3-4 visual Section Panels (panel title + 2-3 bullet insights), and a concise Takeaway footer. Format clearly with labels like [TITLE], [STAT], [PANEL], [TAKEAWAY] so it can be handed to a designer.
   - 'Press Release': Publication-ready release with headline, subhead, dateline, lead, body, quote placeholder, boilerplate, and media contact/CTA placeholders.
   - 'Email Brief': Ready-to-send email with subject line, preheader, concise body, scannable sections, and a clear CTA.
   - 'FAQ': 8-12 anticipated stakeholder questions with accurate, plain-language answers and a next-step section.
3. WRITING STYLE: Use plain professional language, informative headings, short paragraphs, and clean bullets. Do not use emojis, decorative symbols, excessive exclamation marks, or filler. Keep formatting easy to scan and ready to publish.
4. OUTPUT STRICTLY AS A VALID JSON OBJECT where keys are the exact requested format names and values are strings containing the complete, formatted text."""

    fidelity_instruction = (
        "Use only facts supported by the supplied source. Clearly label any unavoidable assumptions."
        if source_fidelity
        else "You may make clearly signposted, reasonable editorial inferences beyond the supplied source."
    )
    user_prompt = f"""Target Formats: {out_list}
Tone: {tone}
Target Audience: {audience}
Output Language: {language}
Detail Level: {detail_level}
Communication Objective: {objective}
Content Style: {style}
Source Fidelity: {fidelity_instruction}

Source Content to Transform:
{combined_text}

Generate the detailed JSON response now:"""
    
    def extract_json(text: str) -> dict:
        """Try to parse JSON even if the model returned markdown wrappers or extra text."""
        text = text.strip()
        # Strip markdown code fences
        text = re.sub(r'^```(?:json)?\s*', '', text)
        text = re.sub(r'\s*```$', '', text)
        text = text.strip()
        try:
            return json.loads(text)
        except Exception:
            # Try to extract the outermost {...} block
            match = re.search(r'(\{[\s\S]*\})', text)
            if match:
                return json.loads(match.group(1))
            raise

    # Model cascade: try each in order, with and without json_object mode
    model_attempts = [
        ("openai/gpt-oss-120b", True),
        ("openai/gpt-oss-20b",  True),
        ("openai/gpt-oss-120b", False),  # plain text fallback — extract JSON manually
        ("openai/gpt-oss-20b",  False),
    ]

    last_error = None
    for model_name, use_json_mode in model_attempts:
        try:
            print(f"[Transform] Trying model={model_name}, json_mode={use_json_mode}")
            kwargs = dict(
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user",   "content": user_prompt}
                ],
                model=model_name,
                temperature=0.7,
            )
            if use_json_mode:
                kwargs["response_format"] = {"type": "json_object"}

            completion = await run_in_threadpool(groq_client.chat.completions.create, **kwargs)
            result = extract_json(completion.choices[0].message.content)
            # Keep generated copy professional and consistent across models.
            emoji_pattern = re.compile(r'[\U0001F000-\U0001FAFF\u2600-\u27BF\uFE0F]')
            result = {key: emoji_pattern.sub('', value) if isinstance(value, str) else value for key, value in result.items()}
            # Validate it has at least one requested key
            if any(k in result for k in out_list):
                print(f"[Transform] Success on model={model_name}")
                try:
                    session_email = _authenticated_email(request)
                    history_record = {
                        "type": "chat_history", "id": str(uuid.uuid4()), "created_at": int(time.time()),
                        "source": combined_text, "outputs": out_list, "tone": tone, "audience": audience,
                        "language": language, "detail_level": detail_level, "objective": objective,
                        "style": style, "source_fidelity": source_fidelity, "result": result,
                    }
                    await run_in_threadpool(_save_exact_history, session_email, history_record)
                    await run_in_threadpool(_save_memory, session_email, history_record)
                except Exception as memory_error:
                    print("[Mem0] history save failed:", memory_error)
                return result
            else:
                raise ValueError(f"Model returned JSON but with unexpected keys: {list(result.keys())}")
        except Exception as exc:
            print(f"[Transform] Failed model={model_name}, json_mode={use_json_mode}: {exc}")
            last_error = exc
            continue

    raise HTTPException(status_code=502, detail="AI generation is temporarily unavailable.")

def _authenticated_email(request: Request):
    token = request.cookies.get("session")
    if not token or "." not in token: return None
    encoded, signature = token.rsplit(".", 1)
    if not hmac.compare_digest(signature, hmac.new((AUTH_SECRET or "dev-secret").encode(), encoded.encode(), hashlib.sha256).hexdigest()): return None
    try:
        email, expiry = base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4)).decode().rsplit(":", 1)
        return email if time.time() < int(expiry) else None
    except (ValueError, TypeError): return None

@app.get("/history")
@app.get("/api/history")
def history(request: Request):
    email = _authenticated_email(request)
    if not email: raise HTTPException(status_code=401, detail="Authentication required.")
    try:
        entries = _read_exact_history(email)
        # Local JSON is the canonical history source; Mem0 is best-effort context storage.
        entries.sort(key=lambda x: x.get("created_at", 0), reverse=True)
        return {"history": entries[:50]}
    except Exception as exc:
        print("[Mem0] history read failed:", exc)
        raise HTTPException(status_code=502, detail="History service unavailable.")

@app.delete("/history/{history_id}")
@app.delete("/api/history/{history_id}")
def delete_history(history_id: str, request: Request):
    email = _authenticated_email(request)
    if not email: raise HTTPException(status_code=401, detail="Authentication required.")
    with history_lock:
        records = json.loads(HISTORY_FILE.read_text(encoding="utf-8")) if HISTORY_FILE.exists() else {}
        entries = records.get(email, [])
        if not any(str(entry.get("id")) == history_id for entry in entries):
            raise HTTPException(status_code=404, detail="History item not found.")

    try:
        _delete_memory_copy(email, history_id)
    except Exception as memory_error:
        print("[Mem0] history deletion failed:", memory_error)
        raise HTTPException(status_code=502, detail="History item could not be deleted from all stores.")

    with history_lock:
        records = json.loads(HISTORY_FILE.read_text(encoding="utf-8")) if HISTORY_FILE.exists() else {}
        entries = records.get(email, [])
        remaining = [entry for entry in entries if str(entry.get("id")) != history_id]
        if len(remaining) == len(entries): raise HTTPException(status_code=404, detail="History item not found.")
        records[email] = remaining
        HISTORY_FILE.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
        return {"status": "deleted"}
